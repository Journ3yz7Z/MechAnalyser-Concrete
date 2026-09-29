"""Current specimen export; intentionally no batch or cross-specimen collation."""
from datetime import datetime
from pathlib import Path
import csv
import re
import uuid
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill
from . import VERSION, UPSTREAM_COMMIT
from .persistence import save_analysis, source_state
from .presentation import injection_text

AXIAL = "#2166ac"
HOOP = "#00846b"
SELECTED = "#de8500"
FIT = "#ae3376"
MODE_LABELS = {"ratio": "分支峰值应力比例", "stress": "轴向应力 MPa", "strain": "轴向应变 mm/mm"}


def summary(a):
    r, p, raw = a.result, a.parameters, a.raw_data
    s = raw.settings
    d = {"源文件": raw.source, "源文件SHA256": raw.fingerprint, "sheet": s.sheet,
         "完整试样编号": a.metadata.get("specimen", ""), "压注压强_MPa": a.metadata.get("pressure", ""),
         "压注时间": injection_text(a.metadata.get("injection_hours", "")), "养护龄期_d": a.metadata.get("curing_days", ""),
         "试验类型": a.metadata.get("test_type") or "单轴压缩试验",
         "围压_MPa": a.metadata.get("confining_pressure", "") if '三轴' in str(a.metadata.get('test_type', '')) else "",
         "材料": a.metadata.get("material", ""), "水泥比例值": a.metadata.get("cement_ratio", ""),
         "细骨料比例值": a.metadata.get("aggregate_ratio", ""), "含水率_%": a.metadata.get("water_content", ""),
         "峰值轴向应力_MPa": r.get("peak_stress_MPa"), "峰值对应轴向应变_mm/mm": r.get("peak_axial_strain"),
         "峰值轴力_kN": r.get("peak_force_kN"), "峰值轴力来源": r.get("peak_force_source", "未提供轴力列或截面积"),
         "轴力峰值原始行": r.get("peak_force_row"), "输入截面积_mm2": s.area_mm2,
         "轴力源列号": s.force_column + 1 if s.force_column is not None else None,
         "轴力源单位": s.force_unit, "轴力符号乘数": s.force_sign,
         "峰值原始行": r.get("peak_row"), "结果名称": "选定区间回归模量",
         "E_MPa": r.get("E", {}).get("slope"), "E_GPa": r.get("E_GPa"),
         "轴向应力—环向应变斜率_kh_MPa": r.get("kh", {}).get("slope"),
         "泊松比_主要结果_斜率比": r.get("nu"), "泊松比_应变回归参考": r.get("nu_strain"),
         "泊松比计算方法": "ASTM D7012-14 公式(7)：ν=-E/k横；k横为轴向应力—横向应变拟合斜率，与E使用相同单位"}
    for key, label in (("E", "轴向"), ("kh", "环向"), ("poisson", "应变回归")):
        fit = r.get(key, {})
        d.update({label + "_斜率": fit.get("slope"), label + "_截距": fit.get("intercept"),
                  label + "_R2": fit.get("r2"), label + "_点数": fit.get("n")})
    d.update({"区间规则": MODE_LABELS[p.mode], "规则下限": p.low, "规则上限": p.high,
              "分支起始原始行": p.branch_start, "分支终止原始行": p.branch_end,
              "分支识别方法": p.branch_method, "分支峰值_MPa": r.get("branch_peak_stress_MPa"),
              "阈值下限": r.get("threshold_low"), "阈值上限": r.get("threshold_high"),
              "同步有效点数": r["n"], "拟合方法": "三项均强制过原点" if p.through_origin else "三项均带截距最小二乘",
              "R2定义": "1-SSE/Σ(y-y平均)²；过原点也使用中心化R²，可为负；恒定y留空",
              "符号约定": s.sign_note, "归一化约定": "轴压正、环胀负；ν=-E/kh，ν_strain=-b（参考）；不取绝对值",
              "环向通道来源": s.channel_source, "检查状态": a.status, "备注": a.notes,
              "源文件状态": a.source_state, "数据问题条数": len(raw.issues),
              "软件版本": VERSION, "原项目提交": UPSTREAM_COMMIT})
    for c, label in (("axial", "轴向应变"), ("hoop", "环向应变"), ("stress", "轴向应力")):
        d[label + "_源列号"] = s.columns[c] + 1
        d[label + "_源单位"] = s.units[c]
        d[label + "_符号乘数"] = s.signs[c]
        bounds = r.get("actual", {}).get(c, [None, None])
        d[label + "_实际下限"], d[label + "_实际上限"] = bounds
    rows = r.get("actual", {}).get("row", [None, None])
    d["拟合最小原始行"], d["拟合最大原始行"] = rows
    d["提示"] = "；".join(r["errors"] + r["warnings"])
    return d


def append_safe(ws, row):
    """Keep user text literal, including strings beginning with '='."""
    ws.append(row)
    for cell, value in zip(ws[ws.max_row], row):
        if isinstance(value, str):
            cell.data_type = "s"


def export_plot(a, path):
    from .figure import export_figure
    return export_figure(a, path)


def export_current(a, parent, formats=("png",)):
    formats = tuple(dict.fromkeys(formats))
    if not formats or any(f not in ("png", "svg") for f in formats):
        raise ValueError("请至少选择一种图片格式（PNG / SVG）。")
    source_state(a)
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", a.metadata.get("specimen") or "当前试样")[:100].rstrip('. ')
    folder = Path(parent) / f'{stem}_{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:6]}'
    folder.mkdir(parents=True, exist_ok=False)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "当前试样结果"
    d = summary(a)
    append_safe(ws, list(d.keys()))
    append_safe(ws, list(d.values()))
    field = wb.create_sheet("结果纵览")
    append_safe(field, ["字段", "值"])
    for k, v in d.items():
        append_safe(field, [k, v])
    selected = wb.create_sheet("实际拟合数据")
    headers = ["原始Excel行号", "轴向应变_mm/mm", "环向应变_mm/mm", "轴向应力_MPa"]
    append_safe(selected, headers)
    selected_rows = a.selected[["row", "axial", "hoop", "stress"]].values.tolist()
    for row in selected_rows:
        append_safe(selected, row)
    source = wb.create_sheet("原始值及转换")
    append_safe(source, ["原始行", "轴向原始值", "环向原始值", "应力原始值", "轴向_mm/mm", "环向_mm/mm", "应力_MPa", "三通道有效", "参与拟合"])
    used = set(a.selected.row)
    for raw, row in zip(a.raw_data.cells, a.raw_data.frame.itertuples()):
        nums = [v if np.isfinite(v) else None for v in (row.axial, row.hoop, row.stress)]
        append_safe(source, [row.row, raw["axial"], raw["hoop"], raw["stress"], *nums, bool(row.valid), row.row in used])
    issues = wb.create_sheet("数据问题")
    if "force_kN" in a.raw_data.frame:
        force_sheet = wb.create_sheet("轴力数据")
        append_safe(force_sheet, ["原始Excel行号", "轴力原始值", "轴力_kN"])
        for raw_cell, row in zip(a.raw_data.cells, a.raw_data.frame.itertuples()):
            append_safe(force_sheet, [row.row, raw_cell.get("force"), row.force_kN if np.isfinite(row.force_kN) else None])
    append_safe(issues, ["原始行", "原因", "同步排除"])
    for issue in a.raw_data.issues:
        append_safe(issues, [issue["row"], issue["reason"], issue["excluded"]])
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244E63")
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(55, max(18, len(str(col[0].value)) * 1.6))
    wb.save(folder / "当前试样结果.xlsx")
    with (folder / "实际拟合数据_Origin.csv").open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(selected_rows)
    from .figure import export_figure
    for fmt in formats:
        export_figure(a, folder / f"全范围图.{fmt}")
        export_figure(a, folder / f"局部放大图.{fmt}", xlim=(-.005,.005))
    save_analysis(a, folder / "分析记录.mca.json")
    (folder / "说明.txt").write_text(
        "本目录仅导出当前一个试样。实际拟合数据按原始顺序包含同一批同步点，可直接在 Origin 复核。\n"
        "全范围图和局部放大图使用同一计算结果及纵轴范围；局部图横轴固定−0.005至0.005，超出范围的曲线仅在显示上裁剪。\n"
        "试验类型选项仅标注试验类型，不自动进行围压/偏应力修正，也不改变回归算法。\n"
        "所有拟合使用同一方法；轴向/环向应力拟合的截距单位为 MPa，应变回归截距为 mm/mm。\n"
        "主要泊松比采用 ASTM D7012-14 公式(7)：ν=-E/k横；k横为轴向应力—横向应变曲线斜率，计算时与E均用MPa。图中公式为符号简写。应变回归值仅作参考。\n"
        "数据未平滑、未重排、未默认取绝对值。源数据及其单位/符号转换均已记录。\n"
        "分析记录含原始快照；重新打开会重算并核对源文件指纹。\n"
        "这是选定区间回归模量，未经标准符合性核验。通道来源待核查时结果也需核查。\n"
        + a.status + "\n" + "；".join(a.result["errors"] + a.result["warnings"]), encoding="utf-8")
    return folder
