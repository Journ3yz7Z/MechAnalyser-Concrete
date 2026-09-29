"""Read exactly one worksheet, keeping cells, Excel rows and units traceable."""
from dataclasses import dataclass, asdict
from pathlib import Path
import hashlib
import math
import re
import numpy as np
import openpyxl
import pandas as pd
from mech_analyser.experiment import data as base
from mech_analyser.util.units import unit_registry

CHANNELS = ("axial", "hoop", "stress")
LABELS = ("轴向应变", "环向应变", "轴向应力")
UNITS = {"axial": ("mm/mm", "%"), "hoop": ("mm/mm", "%"),
         "stress": ("MPa", "kPa", "Pa", "GPa")}
ALIASES = (("轴向应变", "轴应变", "axialstrain", "longitudinalstrain"),
           ("环向应变", "横向应变", "径向应变", "hoopstrain", "lateralstrain", "transversestrain"),
           ("轴向应力", "轴应力", "应力", "axialstress", "compressivestress", "stress"))


def file_hash(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def ensure_dimensions(ws):
    # Streaming writers may omit the optional worksheet dimension element.
    if ws.max_row is None or ws.max_column is None:
        ws.calculate_dimension(force=True)


def workbook_sheets(path):
    """Only worksheet metadata; no other specimen's cells are loaded or analysed."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for s in wb:
            ensure_dimensions(s)
        return [s.title for s in wb if (s.max_row or 0) >= 2 and (s.max_column or 0) >= 3]
    finally:
        wb.close()


def sheet_preview(path, sheet, header_row=1):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        ws = wb[sheet]
        ensure_dimensions(ws)
        rows = list(ws.iter_rows(min_row=header_row, max_row=header_row + 7,
                                max_col=min(ws.max_column, 100), values_only=True))
        return rows, ws.max_row, ws.max_column
    finally:
        wb.close()


def infer_columns(headers):
    columns, units = {}, {}
    for channel, names in zip(CHANNELS, ALIASES):
        found = [i for i, v in enumerate(headers) if any(
            a in re.sub(r"\s+", "", str(v)).lower() for a in names)]
        columns[channel] = found[0] if len(found) == 1 else None
        h = str(headers[found[0]]) if len(found) == 1 else ""
        candidates = [u for u in UNITS[channel] if (
            (u in ("mm/mm", "%") and u in h) or
            (channel == "stress" and re.search(r"(?<![a-zA-Z])" + u + r"(?![a-zA-Z])", h)))]
        units[channel] = candidates[0] if len(candidates) == 1 else ""
    return columns, units


def sheet_identity(sheet):
    m = re.fullmatch(r"\s*(\d+(?:\.\d+)?)\s*MPa_(.+)", sheet, re.I)
    return (m[2], m[1]) if m else (sheet, "")


@dataclass
class ImportSettings:
    sheet: str
    header_row: int
    columns: dict
    units: dict
    signs: dict
    sign_note: str = "保持原符号；轴压正、环胀负约定待核查"
    channel_source: str = "此前整理的横向应变通道；来源/传感器原理待核查"
    force_column: int | None = None
    force_unit: str = ""
    force_sign: int = 1
    area_mm2: float | None = None

    def validate(self):
        if self.force_column is not None:
            if not isinstance(self.force_column, int) or self.force_column < 0 or self.force_column in self.columns.values():
                raise ValueError("轴力列必须与三通道不同")
            if self.force_unit not in ("N", "kN", "MN") or self.force_sign not in (-1, 1):
                raise ValueError("请明确轴力单位及压缩正负号")
        if self.area_mm2 is not None and (not math.isfinite(self.area_mm2) or self.area_mm2 <= 0):
            raise ValueError("截面积必须是大于零的有限数，单位 mm²")
        if self.header_row < 1:
            raise ValueError("表头行必须大于等于 1")
        chosen = [self.columns.get(c) for c in CHANNELS]
        if any(v is None or not isinstance(v, int) or v < 0 for v in chosen):
            raise ValueError("缺失通道：请明确指定轴向应变、环向应变和轴向应力列")
        if len(set(chosen)) != 3:
            raise ValueError("三个通道必须选择三个不同的列")
        for c, label in zip(CHANNELS, LABELS):
            if self.units.get(c) not in UNITS[c]:
                raise ValueError(f"{label}单位未知，请人工选择；不能根据数值猜测")
            if self.signs.get(c) not in (-1, 1):
                raise ValueError(f"{label}必须明确选择保持或反转符号")


class RawTranscoder(base.RawTranscoder):
    """Excel adapter replacing the upstream per-column CSV reader."""
    def __init__(self, settings):
        super().__init__({c: base.Transcoder.Column(name=c, input_name=c, output_name=c,
                         input_header_column=settings.columns[c], output_header_column=i)
                          for i, c in enumerate(CHANNELS)})
        self.settings = settings


class RawData(base.RawData):
    def __init__(self, frame, settings, source, fingerprint, cells, issues, headers):
        super().__init__(frame, RawTranscoder(settings))
        self.settings = settings
        self.source = str(Path(source).resolve())
        self.fingerprint = fingerprint
        self.cells = cells
        self.issues = issues
        self.headers = headers

    @classmethod
    def load_sheet(cls, path, settings):
        settings.validate()
        path = Path(path)
        before = file_hash(path)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        frame, cells, issues = [], [], []
        factors = {c: float(unit_registry.Quantity(1, settings.units[c]).to(
            "dimensionless" if c != "stress" else "MPa").magnitude) * settings.signs[c]
                   for c in CHANNELS}
        try:
            ws = wb[settings.sheet]
            ensure_dimensions(ws)
            if max(settings.columns.values()) >= ws.max_column:
                raise ValueError("所选列超出 sheet 范围")
            if settings.force_column is not None and settings.force_column >= ws.max_column:
                raise ValueError("轴力列超出 sheet 范围")
            headers = [str(v or "") for v in next(ws.iter_rows(
                min_row=settings.header_row, max_row=settings.header_row, values_only=True))]
            for rownum, row in enumerate(ws.iter_rows(min_row=settings.header_row + 1,
                                                      values_only=True), settings.header_row + 1):
                original, converted, reasons = {}, {}, []
                for c, label in zip(CHANNELS, LABELS):
                    v = row[settings.columns[c]]
                    original[c] = v if v is None or isinstance(v, (str, int, float, bool)) else str(v)
                    try:
                        if v is None or isinstance(v, bool) or (isinstance(v, str) and v.startswith("=")):
                            raise ValueError()
                        num = float(v)
                        if not math.isfinite(num):
                            raise ValueError()
                        converted[c] = num * factors[c]
                        if isinstance(v, str):
                            issues.append({"row": rownum, "reason": f"{label}文本数字已显式解析", "excluded": False})
                    except (ValueError, TypeError, OverflowError):
                        converted[c] = np.nan
                        reasons.append(f"{label}空值/非数值/公式未求值")
                if settings.force_column is not None:
                    v = row[settings.force_column]
                    original["force"] = v if v is None or isinstance(v, (str, int, float, bool)) else str(v)
                    try:
                        if v is None or isinstance(v, bool):
                            raise ValueError()
                        force = float(v) * {"N": .001, "kN": 1., "MN": 1000.}[settings.force_unit] * settings.force_sign
                        if not math.isfinite(force):
                            raise ValueError()
                        converted["force_kN"] = force
                    except (ValueError, TypeError, OverflowError):
                        converted["force_kN"] = np.nan
                        issues.append({"row": rownum, "reason": "轴力无效；不影响三通道拟合，峰值仅依据有效轴力", "excluded": False})
                if reasons:
                    issues.append({"row": rownum, "reason": "；".join(reasons), "excluded": True})
                cells.append({"row": rownum, **original})
                frame.append({"row": rownum, **converted, "valid": not reasons})
        finally:
            wb.close()
        if file_hash(path) != before:
            raise ValueError("工作簿在读取时发生变化，请重新导入")
        if not frame:
            raise ValueError("所选 sheet 没有数据行")
        return cls(pd.DataFrame(frame), settings, path, before, cells, issues, headers)

    def snapshot(self):
        records = self.frame.astype(object).where(pd.notna(self.frame), None).to_dict("records")
        return dict(source=self.source, fingerprint=self.fingerprint, settings=asdict(self.settings),
                    cells=self.cells, issues=self.issues, headers=self.headers, frame=records)

    @classmethod
    def from_snapshot(cls, d):
        settings = ImportSettings(**d["settings"])
        settings.validate()
        frame = pd.DataFrame(d["frame"])
        for c in CHANNELS:
            frame[c] = pd.to_numeric(frame[c], errors="coerce")
        if "force_kN" in frame:
            frame["force_kN"] = pd.to_numeric(frame["force_kN"], errors="coerce")
        frame["valid"] = np.isfinite(frame[list(CHANNELS)]).all(axis=1)
        rows = frame["row"].to_numpy()
        if len(rows) == 0 or not np.all(np.diff(rows) > 0):
            raise ValueError("分析记录原始行号无效")
        return cls(frame, settings, d["source"], d["fingerprint"], d["cells"], d["issues"], d["headers"])
