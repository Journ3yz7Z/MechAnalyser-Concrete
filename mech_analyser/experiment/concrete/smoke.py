"""Repeatable desktop and packaged-executable end-to-end check.

Uses actual Qt widgets/mouse events; only OS file pickers are supplied paths.
No test sample is ever presented as measured data.
"""
from pathlib import Path
import argparse
import json
import traceback
import numpy as np
import openpyxl
from PySide6.QtCore import Qt, QTimer, QPointF, QEvent
from PySide6.QtGui import QMouseEvent
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox
from .ui import ImportDialog
from .data import file_hash
from .persistence import load_analysis
from mech_analyser.experiment.registry import Registry
from mech_analyser.config import DEFAULT_DESKTOP_EXPERIMENT


def synthetic(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "5.0MPa_TEST-1350-1"
    ws.append(["轴向应变（X，mm/mm）", "环向应变（X，mm/mm）", "轴向应力（Y，MPa）"])
    for e in np.linspace(0, .002, 201):
        ws.append([float(e), float(-.2*e - .00001), float(30000*e+1)])
    for i, stress in enumerate(np.linspace(60,5,101)):
        ws.append([float(.0021+i*.00002), float(-.0005-i*.00001), float(stress)])
    wb.save(path)
    return ws.title


def run(argv):
    parser = argparse.ArgumentParser()
    parser.add_argument('output')
    parser.add_argument('--workbook')
    parser.add_argument('--sheet')
    args = parser.parse_args(argv)
    out = Path(args.output).resolve()
    out.mkdir(parents=True, exist_ok=True)
    path = Path(args.workbook).resolve() if args.workbook else out/'SYNTHETIC_TEST_ONLY.xlsx'
    sheet = args.sheet if args.workbook else synthetic(path)
    if not sheet:
        raise ValueError('Real-data self-test requires an explicitly selected --sheet')
    app = QApplication.instance() or QApplication([])
    window = Registry.desktop_registry().create_desktop(DEFAULT_DESKTOP_EXPERIMENT)
    window.showMaximized()
    report = {"kind": "real worksheet" if args.workbook else "SYNTHETIC_TEST_ONLY", "sheet": sheet,
              "source": str(path), "qt_platform": app.platformName(),
              "file_picker": "test paths injected; actual widgets and mouse events used", "passed": False}
    errors = []
    QMessageBox.critical = lambda parent,title,text,*a: errors.append(f'{title}: {text}')
    QMessageBox.warning = lambda parent,title,text,*a: report.setdefault('dialogs',[]).append(f'{title}: {text}')
    QMessageBox.information = lambda *a: None
    QFileDialog.getOpenFileName = lambda *a,**k: (str(path), '')
    QFileDialog.getSaveFileName = lambda *a,**k: (str(out/'current.mca.json'), '')
    QFileDialog.getExistingDirectory = lambda *a,**k: str(out/'export')

    def execute():
        try:
            before_hash = file_hash(path)
            def choose_sheet():
                d = QApplication.activeModalWidget()
                if not isinstance(d, ImportDialog):
                    errors.append('Import dialog not active')
                    return
                d.sheet.setCurrentIndex(d.sheet.findData(sheet))
                if hasattr(d, 'choices'):
                    for i in range(d.choices.count()):
                        item = d.choices.item(i)
                        item.setCheckState(Qt.CheckState.Checked if item.text() == sheet else Qt.CheckState.Unchecked)
                d.grab().save(str(out/'import-dialog.png'))
                d.validate_accept()
                if not d.result():
                    errors.append('Import validation failed: '+str(report.get('dialogs',[])))
                    d.reject()
            QTimer.singleShot(150, choose_sheet)
            QTest.mouseClick(window.import_button, Qt.MouseButton.LeftButton)
            app.processEvents()
            assert window.analyser is not None, errors
            a = window.analyser
            assert not a.result['errors'], a.result
            a.metadata.update(injection_hours='24' if args.workbook else '测试',
                              curing_days='14' if args.workbook else '测试')
            a.notes = ('真实工作簿单试样软件验证；20%–40% 为测试区间，未做人工物理验收；环向来源待核查。'
                       if args.workbook else '仅用于软件测试的合成数据，E=30000 MPa，ν=0.2，带截距。')
            window.set_analyser(a)
            QTest.qWait(100)
            window.grab().save(str(out/'before-drag.png'))
            window.tabs.setCurrentIndex(0)
            QTest.qWait(80)
            panel = window.panels[0]
            vb = panel.plot.getPlotItem().vb
            x = float(a.raw_data.frame.axial.quantile(.6))
            low = panel.region.getRegion()[0]
            new_low = a.result['branch_peak_stress_MPa']*.25
            start = panel.plot.mapFromScene(vb.mapViewToScene(QPointF(x,low)))
            end = panel.plot.mapFromScene(vb.mapViewToScene(QPointF(x,new_low)))
            old_rows = a.selected.row.tolist()
            QTest.mouseMove(panel.plot.viewport(), start)
            QTest.qWait(80)
            QTest.mousePress(panel.plot.viewport(), Qt.MouseButton.LeftButton, pos=start)
            for t in np.linspace(.1,1,10):
                pos = start+(end-start)*float(t)
                event = QMouseEvent(QEvent.Type.MouseMove, QPointF(pos),
                    QPointF(panel.plot.viewport().mapToGlobal(pos)), Qt.MouseButton.NoButton,
                    Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
                QApplication.sendEvent(panel.plot.viewport(), event)
                QTest.qWait(20)
            QTest.mouseRelease(panel.plot.viewport(), Qt.MouseButton.LeftButton, pos=end)
            QTest.qWait(60)
            report['mouse_drag_changed_selection'] = old_rows != a.selected.row.tolist()
            assert report['mouse_drag_changed_selection'], 'Mouse drag did not change selected rows'
            # Deterministic, explicitly stated validation interval.
            window.low.setValue(.2)
            window.high.setValue(.4)
            app.processEvents()
            assert not a.result['errors'], a.result
            report['selected_rows'] = a.selected.row.tolist()
            report['parameters'] = dict(mode=a.parameters.mode, low=a.parameters.low, high=a.parameters.high,
                                        branch_start=a.parameters.branch_start, branch_end=a.parameters.branch_end)
            for i in range(4):
                window.tabs.setCurrentIndex(i)
                QTest.qWait(80)
                window.grab().save(str(out/f'view-{i+1}.png'))
                assert len(window.panels[i].curves[1].xData) == len(a.selected)
            window.focus_selection()
            window.tabs.setCurrentIndex(2)
            QTest.qWait(80)
            window.grab().save(str(out/'fit-detail.png'))
            report['three_views_same_points'] = True
            # Independent reference: read original cells again, construct mask, scipy regressions.
            from scipy.stats import linregress
            wb = openpyxl.load_workbook(path,read_only=True,data_only=True)
            rows = np.array(list(wb[sheet].values)[1:],dtype=float)
            wb.close()
            peak_index = int(np.nanargmax(rows[:,2]))
            start_index = int(np.where(rows[:peak_index+1,2] == np.nanmin(rows[:peak_index+1,2]))[0][-1])
            row_ids = np.arange(2,len(rows)+2)
            mask = ((row_ids >= start_index+2) & (row_ids <= peak_index+2) &
                    np.isfinite(rows).all(axis=1) & (rows[:,2] >= .2*rows[peak_index,2]) & (rows[:,2] <= .4*rows[peak_index,2]))
            assert row_ids[mask].tolist() == a.selected.row.tolist()
            selected = rows[mask]
            expected = {}
            for key, x_col, y_col in [('E',0,2),('kh',1,2),('poisson',0,1)]:
                fit = linregress(selected[:,x_col],selected[:,y_col])
                expected[key] = dict(slope=float(fit.slope),intercept=float(fit.intercept),r2=float(fit.rvalue**2))
                for param in ['slope','intercept','r2']:
                    assert np.isclose(expected[key][param], a.result[key][param],rtol=1e-10,atol=1e-12), (key,param)
            report['independent_scipy'] = expected
            report['result'] = a.result
            QTest.mouseClick(window.save_button, Qt.MouseButton.LeftButton)
            assert (out/'current.mca.json').exists(), errors
            restored = load_analysis(out/'current.mca.json')
            assert restored.result == a.result
            report['restore_identical'] = True
            QTest.mouseClick(window.export_button, Qt.MouseButton.LeftButton)
            exported = list((out/'export').glob('*/当前试样结果.xlsx'))
            assert len(exported) == 1, errors
            assert {p.name for p in exported[0].parent.glob('*.png')} == {'全范围图.png','局部放大图.png'}
            export_wb = openpyxl.load_workbook(exported[0],data_only=True)
            data = list(export_wb['实际拟合数据'].values)[1:]
            assert [row[0] for row in data] == a.selected.row.tolist()
            assert np.allclose(np.array(data)[:,1:], selected,rtol=1e-12,atol=1e-15)
            export_wb.close()
            report['export'] = str(exported[0].parent)
            import xml.etree.ElementTree as ET
            report['format_checks'] = {}
            for formats in (('svg',), ('png', 'svg')):
                before = set((out/'export').iterdir())
                for fmt, check in window.export_formats.items():
                    check.setChecked(fmt in formats)
                QTest.mouseClick(window.export_button, Qt.MouseButton.LeftButton)
                created = set((out/'export').iterdir()) - before
                assert len(created) == 1, errors
                folder = created.pop()
                for fmt in ('png', 'svg'):
                    names = {p.name for p in folder.glob('*.'+fmt)}
                    expected_names = {f'全范围图.{fmt}', f'局部放大图.{fmt}'} if fmt in formats else set()
                    assert names == expected_names, (formats, names, errors)
                for svg in folder.glob('*.svg'):
                    assert ET.parse(svg).getroot().tag.endswith('svg')
                report['format_checks']['+'.join(formats)] = True
            report['source_sha256_before'] = before_hash
            report['source_sha256_after'] = file_hash(path)
            assert before_hash == file_hash(path)
            assert not errors, errors
            report['passed'] = True
        except Exception:
            report['exception'] = traceback.format_exc()
        finally:
            report['ui_errors'] = errors
            (out/'smoke-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
            window.dirty = False
            window.close()
            app.quit()
    QTimer.singleShot(100,execute)
    app.exec()
    print(json.dumps(report,ensure_ascii=True,indent=2))
    return 0 if report['passed'] else 1
