from dataclasses import replace
import json
import numpy as np
import openpyxl
import pytest
from mech_analyser.experiment.concrete.analyser import Analyser, regression
from mech_analyser.experiment.concrete.data import RawData, infer_columns, workbook_sheets, file_hash
from mech_analyser.experiment.concrete.persistence import save_analysis, load_analysis
from mech_analyser.experiment.concrete.collation import export_current


def test_known_E_nu_intercepts(analysis):
    r = analysis.result
    assert not r['errors']
    assert r['E']['slope'] == pytest.approx(30000)
    assert r['E']['intercept'] == pytest.approx(2)
    assert r['nu'] == pytest.approx(.2)
    assert r['kh']['slope'] == pytest.approx(-150000)
    assert r['poisson']['intercept'] == pytest.approx(-.00001)
    assert r['E']['r2'] == pytest.approx(1)
    assert r['nu_ratio'] == pytest.approx(.2)


@pytest.mark.parametrize('axial_unit,hoop_unit,stress_unit', [('mm/mm','mm/mm','MPa'),('%','%','kPa'),('%','mm/mm','Pa'),('mm/mm','%','GPa')])
def test_units_equal(workbook, settings, axial_unit, hoop_unit, stress_unit):
    factors = {'mm/mm': 1, '%': 100, 'MPa': 1, 'kPa': 1000, 'Pa': 1e6, 'GPa': .001}
    units = [axial_unit, hoop_unit, stress_unit]
    wb = openpyxl.load_workbook(workbook)
    for row in wb.active.iter_rows(min_row=2):
        for cell, unit in zip(row, units):
            cell.value *= factors[unit]
    wb.save(workbook)
    settings.units = dict(zip(('axial','hoop','stress'), units))
    a = Analyser(RawData.load_sheet(workbook, settings))
    assert a.result['E']['slope'] == pytest.approx(30000)
    assert a.result['nu'] == pytest.approx(.2)


def test_origin_independent_matrix_solution(analysis):
    analysis.parameters.through_origin = True
    r = analysis.analyse()
    s = analysis.selected
    expected = np.linalg.lstsq(s.axial.to_numpy()[:, None], s.stress, rcond=None)[0][0]
    assert r['E']['slope'] == pytest.approx(expected)
    assert r['E']['intercept'] == 0
    assert r['E']['slope'] != pytest.approx(30000)
    expected_b = np.linalg.lstsq(s.axial.to_numpy()[:, None], s.hoop, rcond=None)[0][0]
    assert r['nu_strain'] == pytest.approx(-expected_b)
    expected_k = np.linalg.lstsq(s.hoop.to_numpy()[:, None], s.stress, rcond=None)[0][0]
    assert r['nu'] == pytest.approx(-expected / expected_k)


@pytest.mark.parametrize('mode,low,high', [('ratio',.2,.4), ('stress',25,49), ('strain',.0007,.0015)])
def test_postpeak_never_mixed(analysis, mode, low, high):
    p = analysis.parameters
    p.mode, p.low, p.high = mode, low, high
    analysis.analyse()
    assert analysis.selected.row.max() < 122
    assert len(analysis.selected) > 3
    assert analysis.result['E']['slope'] == pytest.approx(30000)
    assert analysis.parameters.branch_end == 122


def test_missing_values_keep_rows_synchronous(workbook, settings):
    wb = openpyxl.load_workbook(workbook)
    wb.active.cell(30, 2).value = None
    wb.active.cell(31, 1).value = 'bad'
    wb.active.cell(32, 3).value = '=1+2'
    wb.save(workbook)
    a = Analyser(RawData.load_sheet(workbook, settings))
    assert not {30,31,32}.intersection(a.selected.row)
    assert len(a.raw_data.frame) == 221
    assert len(a.raw_data.issues) == 3
    for row in a.selected.itertuples():
        assert row.stress == pytest.approx(2 + 30000 * row.axial)
        assert row.hoop == pytest.approx(-.00001 - .2 * row.axial)
    assert all(a.result[k]['n'] == len(a.selected) for k in ['E','kh','poisson'])


@pytest.mark.parametrize('low,high', [(2,3),(.4,.2),(.4,.4),(-.1,.5)])
def test_invalid_range_no_fake_values(analysis, low, high):
    analysis.parameters.low, analysis.parameters.high = low, high
    r = analysis.analyse()
    assert r['status'] == '计算失败'
    assert r['errors']
    assert 'E' not in r and 'nu' not in r


@pytest.mark.parametrize('values', [([1,1,1],[2,3,4]),([0,1],[1,2]),([0,float('nan'),2],[1,2,3])])
def test_regression_rejects_bad_points(values):
    with pytest.raises(ValueError):
        regression(*values)


def test_no_absolute_values(analysis):
    analysis.raw_data.frame['hoop'] *= -1
    r = analysis.analyse()
    assert r['nu'] < 0
    assert r['kh']['slope'] > 0
    assert any('符号' in w for w in r['warnings'])


def test_missing_channel_fails(workbook, settings):
    settings.columns['hoop'] = None
    with pytest.raises(ValueError, match='缺失通道'):
        RawData.load_sheet(workbook, settings)


def test_unknown_units_not_guessed(workbook, settings):
    settings.units['hoop'] = ''
    with pytest.raises(ValueError, match='单位未知'):
        RawData.load_sheet(workbook, settings)
    cols, units = infer_columns(['轴向应变','环向应变','轴向应力'])
    assert cols == dict(axial=0,hoop=1,stress=2)
    assert all(not v for v in units.values())


def test_manual_column_permutation(workbook, settings):
    wb = openpyxl.load_workbook(workbook)
    for row in wb.active:
        row[0].value, row[2].value = row[2].value, row[0].value
    wb.save(workbook)
    settings.columns = dict(axial=2,hoop=1,stress=0)
    a = Analyser(RawData.load_sheet(workbook, settings))
    assert a.result['E']['slope'] == pytest.approx(30000)


def test_zero_hoop_slope_no_infinity(analysis):
    analysis.parameters.mode = 'strain'
    analysis.parameters.low, analysis.parameters.high = 0, .004
    analysis.raw_data.frame['stress'] = 20.
    r = analysis.analyse()
    assert r['nu_ratio'] is None
    assert r['nu'] is None
    assert r['kh']['r2'] is None
    assert any('接近零' in w for w in r['warnings'])


def test_noisy_nu_methods_differ(analysis):
    rng = np.random.default_rng(1701)
    analysis.raw_data.frame['hoop'] += rng.normal(0, .00005, len(analysis.raw_data.frame))
    r = analysis.analyse()
    assert abs(r['nu_strain'] - r['nu']) > .001
    s = analysis.selected
    e = np.linalg.lstsq(np.column_stack([s.axial, np.ones(len(s))]), s.stress, rcond=None)[0][0]
    k = np.linalg.lstsq(np.column_stack([s.hoop, np.ones(len(s))]), s.stress, rcond=None)[0][0]
    assert r['nu'] == pytest.approx(-e/k)


def test_explicit_sign_conversion(workbook, settings):
    wb = openpyxl.load_workbook(workbook)
    for row in wb.active.iter_rows(min_row=2):
        for cell in row:
            cell.value *= -1
    wb.save(workbook)
    settings.signs = dict(axial=-1,hoop=-1,stress=-1)
    a = Analyser(RawData.load_sheet(workbook, settings))
    assert a.result['E']['slope'] == pytest.approx(30000)
    assert a.result['nu'] == pytest.approx(.2)
    assert a.raw_data.cells[1]['axial'] < 0


def test_peak_row_preserved_when_hoop_missing(workbook, settings):
    wb = openpyxl.load_workbook(workbook)
    wb.active.cell(122,2).value = None
    wb.save(workbook)
    a = Analyser(RawData.load_sheet(workbook, settings))
    assert a.parameters.branch_end == 122
    assert a.result['peak_row'] == 122


def test_single_sheet_only(workbook, settings):
    wb = openpyxl.load_workbook(workbook)
    ws = wb.create_sheet('DO_NOT_ANALYSE')
    ws.append(['x','y','z'])
    ws.append(['bad','bad','bad'])
    wb.save(workbook)
    assert len(workbook_sheets(workbook)) == 2
    a = Analyser(RawData.load_sheet(workbook, settings))
    assert not a.result['errors'] and not a.raw_data.issues
    assert len(a.raw_data.frame) == 221


def test_roundtrip_and_source_change(analysis, tmp_path, workbook):
    path = tmp_path / 'current.mca.json'
    analysis.notes = '待核查，仅测试'
    analysis.reviewed = True
    analysis.save(path)
    restored = load_analysis(path)
    assert restored.result == analysis.result
    assert restored.selected.row.tolist() == analysis.selected.row.tolist()
    assert restored.reviewed
    assert restored.notes == analysis.notes
    wb = openpyxl.load_workbook(workbook)
    wb.active.cell(3,3).value = 999
    wb.save(workbook)
    stale = load_analysis(path)
    assert '改变' in stale.source_state
    assert not stale.reviewed
    assert stale.result == analysis.result


def test_tampered_record_rejected(analysis, tmp_path):
    path = tmp_path/'current.mca.json'
    save_analysis(analysis, path)
    d = json.loads(path.read_text(encoding='utf-8'))
    d['parameters']['low'] = .1
    path.write_text(json.dumps(d),encoding='utf-8')
    with pytest.raises(ValueError,match='校验失败'):
        load_analysis(path)


def test_export_source_unchanged_and_exact_rows(analysis, tmp_path, workbook):
    before = file_hash(workbook)
    analysis.notes = '=HYPERLINK("https://invalid.example")'
    folder = export_current(analysis, tmp_path)
    assert file_hash(workbook) == before
    wb = openpyxl.load_workbook(folder/'当前试样结果.xlsx', data_only=False)
    rows = list(wb['实际拟合数据'].values)[1:]
    assert [v[0] for v in rows] == analysis.selected.row.tolist()
    assert np.allclose(np.array(rows)[:,1:], analysis.selected[['axial','hoop','stress']])
    summary = wb['当前试样结果']
    col = list(next(summary.values)).index('备注') + 1
    assert summary.cell(2,col).data_type == 's'
    assert (folder/'全范围图.png').stat().st_size > 10000
    assert (folder/'局部放大图.png').stat().st_size > 10000
    assert load_analysis(folder/'分析记录.mca.json').result == analysis.result
    with pytest.raises(ValueError,match='禁止覆盖'):
        save_analysis(analysis,workbook)


def test_manual_branch_and_review_invalidation(analysis):
    analysis.reviewed = True
    analysis.parameters.branch_start = 35
    analysis.parameters.branch_end = 48
    analysis.parameters.mode = 'strain'
    analysis.parameters.low = 0
    analysis.parameters.high = 1
    analysis.analyse()
    assert analysis.selected.row.tolist() == list(range(35,49))
    assert not analysis.reviewed


def test_default_branch_excludes_prior_unload_reload():
    import pandas as pd
    from mech_analyser.experiment.concrete.phase import default_branch
    f = pd.DataFrame(dict(row=range(2,13),stress=[0,10,20,30,20,10,12,18,25,35,40]))
    assert default_branch(f) == (7,12)
