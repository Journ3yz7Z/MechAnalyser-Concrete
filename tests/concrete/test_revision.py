from dataclasses import replace
import numpy as np
import openpyxl
import pytest
from PySide6.QtWidgets import QDialog, QMessageBox
from PySide6.QtCore import Qt
from mech_analyser.experiment.concrete.data import RawData
from mech_analyser.experiment.concrete.analyser import Analyser
from mech_analyser.experiment.concrete.ui import ConcreteWindow, WorkbookImportDialog, ImportDialog
from mech_analyser.experiment.concrete.persistence import save_workspace, load_workspace, save_analysis, load_analysis
from mech_analyser.experiment.concrete.collation import summary


def test_force_units_priority_and_restore(workbook, settings, tmp_path):
    wb = openpyxl.load_workbook(workbook)
    ws = wb.active
    ws.cell(1, 4, '轴力 (N)')
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 4, -row * 1000)
    ws.cell(4, 4, 'bad')
    wb.save(workbook)
    settings.force_column, settings.force_unit, settings.force_sign = 3, 'N', -1
    settings.area_mm2 = 10000
    a = Analyser(RawData.load_sheet(workbook, settings))
    assert a.result['peak_force_kN'] == ws.max_row
    assert a.result['peak_force_row'] == ws.max_row
    assert '实测' in a.result['peak_force_source']
    path = tmp_path / 'force.mca.json'
    save_analysis(a, path)
    restored = load_analysis(path)
    assert restored.result['peak_force_kN'] == ws.max_row
    assert summary(a)['轴力源单位'] == 'N'


def test_area_and_missing_force(workbook, settings):
    a = Analyser(RawData.load_sheet(workbook, settings))
    assert 'peak_force_kN' not in a.result
    settings.area_mm2 = 10000
    a = Analyser(RawData.load_sheet(workbook, settings))
    assert a.result['peak_force_kN'] == pytest.approx(1220)
    settings.area_mm2 = -1
    with pytest.raises(ValueError, match='截面积'):
        settings.validate()


def test_multi_import_and_independent_progress(workbook, settings, tmp_path, qtbot, monkeypatch):
    wb = openpyxl.load_workbook(workbook)
    second = wb.copy_worksheet(wb.active)
    second.title = 'second'
    wb.save(workbook)
    dialog = WorkbookImportDialog(workbook)
    qtbot.addWidget(dialog)
    assert dialog.choices.parentWidget() is not dialog
    dialog.show()
    qtbot.wait(30)
    assert dialog.choices.isVisible()
    dialog.validate_accept()
    assert dialog.result() == QDialog.Accepted
    assert [s.sheet for s in dialog.all_settings] == [settings.sheet, 'second']
    w = ConcreteWindow()
    qtbot.addWidget(w, before_close_func=lambda x: (x.sessions.clear(), setattr(x, 'dirty', False)))
    w.install_sessions([Analyser(RawData.load_sheet(workbook, s)) for s in dialog.all_settings])
    assert w.result_text.isHidden()
    assert set(w.metrics) == {'stress', 'force', 'E', 'nu'}
    w.more_button.click()
    assert not w.result_text.isHidden()
    w.low.setValue(.25)
    w.notes.setPlainText('first note')
    w.specimens.setCurrentIndex(1)
    assert w.low.value() == .2
    w.low.setValue(.3)
    w.specimens.setCurrentIndex(0)
    assert w.low.value() == .25
    assert w.notes.toPlainText() == 'first note'
    w.stash_session()
    path = tmp_path / 'all.mcw.json'
    save_workspace([s['analyser'] for s in w.sessions], path, 1)
    restored, active = load_workspace(path)
    assert active == 1
    assert [a.parameters.low for a in restored] == [.25, .3]
    assert restored[0].notes == 'first note'
    assert np.array_equal(restored[0].selected.row, w.analyser.selected.row)
    w.sessions.clear()
    w.dirty = False


def test_multi_reject_mismatched_headers(workbook, monkeypatch, qtbot):
    wb = openpyxl.load_workbook(workbook)
    second = wb.copy_worksheet(wb.active)
    second.title = 'wrong units'
    second.cell(1, 1, '轴向应变 (%)')
    wb.save(workbook)
    messages = []
    monkeypatch.setattr(QMessageBox, 'warning', lambda *args: messages.append(args[2]))
    d = WorkbookImportDialog(workbook)
    qtbot.addWidget(d)
    d.validate_accept()
    assert d.result() != QDialog.Accepted
    assert 'wrong units' in messages[0]
    d.choices.item(1).setCheckState(Qt.Unchecked)
    d.validate_accept()
    assert len(d.all_settings) == 1


def test_multi_ignores_only_trailing_blank_headers(workbook, monkeypatch, qtbot):
    wb = openpyxl.load_workbook(workbook)
    second = wb.copy_worksheet(wb.active)
    second.title = 'extra blanks'
    second.cell(1, 9, '   ')
    wb.save(workbook)
    messages=[]
    monkeypatch.setattr(QMessageBox,'warning',lambda *args:messages.append(args[2]))
    d=WorkbookImportDialog(workbook)
    qtbot.addWidget(d)
    d.validate_accept()
    assert d.result() == QDialog.Accepted
    assert len(d.all_settings) == 2
    assert not messages
    second.cell(1, 2).value=None
    wb.save(workbook)
    d2=WorkbookImportDialog(workbook)
    qtbot.addWidget(d2)
    d2.validate_accept()
    assert d2.result() != QDialog.Accepted
    assert messages
