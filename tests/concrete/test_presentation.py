import numpy as np
import pytest
from mech_analyser.experiment.concrete.ui import ConcreteWindow
from mech_analyser.experiment.concrete.presentation import poisson_html, band
from mech_analyser.experiment.concrete.figure import export_figure
from mech_analyser.experiment.concrete.collation import summary
from mech_analyser.experiment.concrete.persistence import save_analysis, load_analysis
from mech_analyser.experiment.concrete.analyser import Analyser


def test_metadata_roundtrip_and_independent_sessions(qtbot, analysis, tmp_path):
    other = Analyser(analysis.raw_data,metadata=analysis.metadata.copy())
    w = ConcreteWindow()
    qtbot.addWidget(w, before_close_func=lambda widget: setattr(widget, 'dirty',False))
    w.install_sessions([analysis,other])
    for key,value in dict(test_type='三轴压缩试验',material='水泥砂浆',cement_ratio='1',aggregate_ratio='2',water_content='13',curing_days='14').items():
        w.fields[key].setText(value)
    w.metadata_changed()
    w.specimens.setCurrentIndex(1)
    assert w.fields['material'].text() == '水泥砂浆'
    assert w.fields['test_type'].text() == '单轴压缩试验'
    w.specimens.setCurrentIndex(0)
    assert w.fields['material'].text() == '水泥砂浆'
    path=tmp_path/'saved.mca.json'
    save_analysis(analysis,path)
    restored=load_analysis(path)
    assert restored.metadata == analysis.metadata
    assert summary(restored)['含水率_%'] == '13'
    assert summary(restored)['试验类型'] == '三轴压缩试验'
    w.sessions.clear()
    w.dirty=False


def test_two_exports_same_results(analysis,tmp_path):
    from mech_analyser.experiment.concrete.collation import export_current
    import copy
    before=copy.deepcopy(analysis.result)
    folder=export_current(analysis,tmp_path)
    assert {p.name for p in folder.glob('*.png')} == {'全范围图.png','局部放大图.png'}
    assert analysis.result == before
    full=export_figure(analysis,tmp_path/'full.png')
    local=export_figure(analysis,tmp_path/'local.png',xlim=(-.005,.005))
    assert local.axes[0].get_xlim() == (-.005,.005)
    assert local.axes[0].get_ylim() == full.axes[0].get_ylim()


def test_local_peak_outside(analysis,tmp_path):
    analysis.result['peak_axial_strain']=.02
    fig=export_figure(analysis,tmp_path/'outside.png',xlim=(-.005,.005))
    ax=fig.axes[0]
    assert ax.get_xlim() == (-.005,.005)
    assert not any(line.get_marker()=='o' for line in ax.lines)
    assert any('峰值位于显示范围外' in t.get_text() for t in ax.texts+fig.texts)


def test_arrow_follows_range_and_formula_preserves_algorithm(qtbot, analysis):
    w=ConcreteWindow()
    qtbot.addWidget(w,before_close_func=lambda widget:setattr(widget,'dirty',False))
    w.set_analyser(analysis)
    w.show()
    qtbot.wait(50)
    w.low.setValue(.27)
    w.high.setValue(.43)
    p=w.panels[3]
    assert np.allclose(p.arrow_line.yData,band(analysis))
    assert '<i>E</i>' in poisson_html(analysis) and '<i>k</i>' in poisson_html(analysis)
    assert 'dε' not in poisson_html(analysis)
    assert f'{analysis.result["nu"]:.4f}' in poisson_html(analysis)
    assert p.zero_axis.isVisible()
    assert analysis.result['nu'] == pytest.approx(.2)
    w.dirty=False


def test_export_single_overlay(analysis,tmp_path):
    path=tmp_path/'overlay.png'
    fig=export_figure(analysis,path)
    assert len(fig.axes) == 1
    ax=fig.axes[0]
    assert ax.get_xlabel() == '应变'
    assert ax.get_ylabel() == '应力(MPa)'
    assert ax.spines['left'].get_position() == ('data',0)
    assert [t.get_text() for t in ax.get_legend().texts] == ['轴向应变','横向应变','拟合曲线']
    assert path.stat().st_size > 10000
    assert np.allclose(fig.get_size_inches()*25.4,[170,103])
    assert ax.xaxis.label.get_size() == 10
    assert ax.spines['bottom'].get_linewidth() == 1
    assert len(ax.collections) == 0
    assert [line.get_linewidth() for line in ax.lines[:4]] == [1.3]*4
    for line in (ax.lines[1],ax.lines[3]):
        assert line.get_color() == '#c62828'
        assert line.get_linestyle() == '-'


def test_streaming_excel_without_dimensions(workbook,settings,tmp_path):
    import openpyxl
    from mech_analyser.experiment.concrete.data import workbook_sheets,sheet_preview,RawData
    source=openpyxl.load_workbook(workbook,read_only=True)
    target=openpyxl.Workbook(write_only=True)
    ws=target.create_sheet(settings.sheet)
    for row in source.active.values:
        ws.append(row)
    source.close()
    path=tmp_path/'streaming.xlsx'
    target.save(path)
    assert workbook_sheets(path) == [settings.sheet]
    assert sheet_preview(path,settings.sheet)[2] == 3
    loaded=RawData.load_sheet(path,settings)
    expected=RawData.load_sheet(workbook,settings)
    assert loaded.frame.equals(expected.frame)


def test_layout_fallback_preserves_visible_results(analysis,tmp_path,monkeypatch):
    import mech_analyser.experiment.concrete.figure as module
    monkeypatch.setattr(module,'clear_box',lambda *args,**kwargs:None)
    fig=module.export_figure(analysis,tmp_path/'fallback.png')
    ax=fig.axes[0]
    result=next(x for x in ax.artists if x.get_gid()=='fit-results')
    rect=result.get_window_extent(fig.canvas.get_renderer())
    assert rect.y1 < ax.bbox.y0
    assert rect.x0>=0 and rect.y0>=0 and rect.x1<fig.bbox.x1


def test_clear_box_refuses_fully_occupied_view(analysis):
    from mech_analyser.experiment.concrete.presentation import clear_box
    assert clear_box(analysis,((-.01,.02),(0,150)),(.2,.1),(.4,.3),[(0,0,1,1)]) is None


def test_overlay_legend_and_resize_axis(qtbot,analysis):
    from PySide6.QtCore import QPointF
    w=ConcreteWindow()
    qtbot.addWidget(w,before_close_func=lambda widget:setattr(widget,'dirty',False))
    w.set_analyser(analysis)
    w.show()
    w.resize(1500,950)
    qtbot.wait(80)
    panel=w.panels[3]
    assert len(panel.plot.getPlotItem().legend.items)==3
    assert not panel.curves[1].isVisible()
    assert not panel.arrow_line.isVisible()
    vb=panel.plot.getViewBox()
    expected=vb.mapFromView(QPointF(0,vb.viewRange()[1][1])).x()
    assert panel.zero_axis.pos().x()+55 == pytest.approx(expected,abs=2)
    w.dirty=False
